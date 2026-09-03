from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth import get_user_model
from django.db.models import Q, Count, Sum, Avg 
from django.db.models.functions import TruncDate, TruncMonth, TruncYear
from django.core.paginator import Paginator
from django.utils import timezone
from django.contrib import messages
from django.http import HttpResponse
from datetime import timedelta, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from orders.models import Order, OrderItem, Coupon
from product.models import Product, ProductImage
from product.form import ProductForm
from category.models import Category
from category.forms import CategoryForm
from category.utils import get_category_tree

import json


User = get_user_model()

def dashboard(request):

    total_revenue = (
        Order.objects.filter(
            payment_status=True,
            status="delivered"
        ).aggregate(
            total=Sum("total")
        )["total"] or 0
    )

    total_orders = Order.objects.count()

    total_customers = User.objects.filter(is_staff=False).count()

    total_products = Product.objects.count()

    recent_orders = Order.objects.order_by("-created_at")[:5]

    context = {
        "total_revenue": total_revenue,
        "total_orders": total_orders,
        "total_customers": total_customers,
        "total_products": total_products,
        "recent_orders": recent_orders,
    }

    return render(request, "dashboard/dashboard.html", context)

def order_list(request):

    search = request.GET.get("search", "")

    status = request.GET.get("status", "")

    payment = request.GET.get("payment", "")

    date_filter = request.GET.get("date", "")

    orders = (
        Order.objects
        .select_related("user")
        .prefetch_related("items")
        .order_by("-created_at")
    )

    if search:
        
        orders = orders.filter(
            Q(user__username__icontains=search) |
            Q(order_number__icontains=search) 
        )

    if status:
        orders = orders.filter(status=status)

    if payment:
        orders = orders.filter(payment_method=payment)

    today = timezone.now().date()

    if date_filter == "today":

        orders = orders.filter(
            created_at__date=today
        )

    elif date_filter == "week":

        orders = orders.filter(
            created_at__date__gte=today - timedelta(days=7)
        )

    elif date_filter == "month":

        orders = orders.filter(
            created_at__month=today.month,
            created_at__year=today.year,
        )

    elif date_filter == "year":

        orders = orders.filter(
            created_at__year=today.year
        )

    paginator = Paginator(orders, 10)

    page = request.GET.get("page")

    orders = paginator.get_page(page)

    total_orders = Order.objects.count()

    pending_orders = Order.objects.filter(
        status="pending"
    ).count()

    delivered_orders = Order.objects.filter(
        status="delivered"
    ).count()

    total_revenue = (
        Order.objects.filter(
            payment_status=True,
            status="delivered"
        ).aggregate(
            total=Sum("total")
        )["total"] or 0
    )

    return render(
        request,
        "dashboard/orders/order_list.html",
        {
            "orders": orders,
            "search": search,
            "status": status,
            "payment": payment,
            "date_filter": date_filter,
            "total_orders": total_orders,
            "pending_orders": pending_orders,
            "delivered_orders": delivered_orders,
            "total_revenue": total_revenue,
        }
    )

def export_orders_excel(request):

    search = request.GET.get("search", "")
    status = request.GET.get("status", "")
    payment = request.GET.get("payment", "")
    date_filter = request.GET.get("date", "")

    orders = (
        Order.objects
        .select_related("user")
        .prefetch_related("items")
        .order_by("-created_at")
    )

    if search:

        orders = orders.filter(
            Q(user__username__icontains=search) |
            Q(order_number__icontains=search)
        )

    if status:

        orders = orders.filter(status=status)

    if payment:

        orders = orders.filter(payment_method=payment)

    today = timezone.now().date()

    if date_filter == "today":

        orders = orders.filter(
            created_at__date=today
        )

    elif date_filter == "week":

        orders = orders.filter(
            created_at__date__gte=today - timedelta(days=7)
        )

    elif date_filter == "month":

        orders = orders.filter(
            created_at__month=today.month,
            created_at__year=today.year
        )

    elif date_filter == "year":

        orders = orders.filter(
            created_at__year=today.year
        )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Orders"

    headers = [

        "Order Number",
        "Customer",
        "Items",
        "Date",
        "Payment Method",
        "Payment Status",
        "Order Status",
        "Total"

    ]

    for col, header in enumerate(headers, start=1):

        cell = sheet.cell(row=1, column=col)

        cell.value = header

        cell.font = Font(bold=True)

    row = 2

    for order in orders:

        sheet.cell(row=row, column=1).value = order.order_number

        sheet.cell(row=row, column=2).value = order.user.username

        sheet.cell(row=row, column=3).value = order.items.count()

        sheet.cell(
            row=row,
            column=4
        ).value = order.created_at.strftime("%d-%m-%Y")

        sheet.cell(
            row=row,
            column=5
        ).value = order.get_payment_method_display()

        sheet.cell(
            row=row,
            column=6
        ).value = "Paid" if order.payment_status else "Pending"

        sheet.cell(
            row=row,
            column=7
        ).value = order.get_status_display()

        sheet.cell(
            row=row,
            column=8
        ).value = float(order.total)

        row += 1

    for column_cells in sheet.columns:

        length = max(
            len(str(cell.value))
            if cell.value else 0
            for cell in column_cells
        )

        sheet.column_dimensions[
            column_cells[0].column_letter
        ].width = length + 5

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    filename = datetime.now().strftime("Orders_%d-%m-%Y.xlsx") + "_Orders.xlsx"

    response[
        "Content-Disposition"
    ] = f'attachment; filename="{filename}"'

    workbook.save(response)

    return response

def order_detail(request, pk):

    order = get_object_or_404(
        Order.objects.select_related(
            "user",
            "address"
        ).prefetch_related("items__product"),
        pk=pk
    )

    context = {
        "order": order,
    }

    return render(
        request,
        "dashboard/orders/order_detail.html",
        context,
    )

def update_order_status(request, pk):

    order = get_object_or_404(Order, pk=pk)

    if request.method == "POST":

        status = request.POST.get("status")

        if status in dict(Order.STATUS_CHOICES):

            order.status = status

            if status == "cancelled":
                order.cancelled_at = timezone.now()

            else:
                order.cancelled_at = None

            order.save()

    return redirect("dashboard:order_detail", pk=pk)

def order_invoice(request, pk):

    order = get_object_or_404(
        Order.objects.select_related(
            "user",
            "address"
        ).prefetch_related(
            "items__product"
        ),
        pk=pk
    )

    return render(
        request,
        "dashboard/orders/invoice.html",
        {
            "order": order,
        }
    )

def download_invoice_pdf(request, pk):

    order = get_object_or_404(
        Order.objects.select_related(
            "user",
            "address"
        ).prefetch_related(
            "items__product"
        ),
        pk=pk
    )

    response = HttpResponse(content_type="application/pdf")

    response["Content-Disposition"] = (
        f'attachment; filename="Invoice-{order.order_number}.pdf"'
    )

    doc = SimpleDocTemplate(response)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph("<b>CASAIO</b>", styles["Title"])
    )

    elements.append(
        Paragraph(
            f"Invoice No: {order.order_number}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"Date: {order.created_at.strftime('%d %b %Y')}",
            styles["Normal"]
        )
    )

    elements.append(Spacer(1, 20))

    elements.append(
        Paragraph(
            "<b>Bill To</b>",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            order.address.full_name,
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            order.address.house_name,
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            order.address.area,
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"{order.address.city}, {order.address.state}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            order.address.pincode,
            styles["Normal"]
        )
    )

    elements.append(Spacer(1, 20))

    data = [
        [
            "Product",
            "Qty",
            "Price",
            "Total",
        ]
    ]

    for item in order.items.all():

        data.append(
            [
                item.product.name,
                str(item.quantity),
                f"₹{item.price}",
                f"₹{item.total_price}",
            ]
        )

    table = Table(data)

    table.setStyle(

        TableStyle(

            [

                ("BACKGROUND", (0,0), (-1,0), colors.orange),

                ("TEXTCOLOR", (0,0), (-1,0), colors.white),

                ("GRID", (0,0), (-1,-1), 1, colors.grey),

                ("BOTTOMPADDING", (0,0), (-1,0), 12),

                ("BACKGROUND", (0,1), (-1,-1), colors.beige),

            ]

        )

    )

    elements.append(table)

    elements.append(Spacer(1, 25))

    elements.append(

        Paragraph(

            f"<b>Total : ₹{order.total}</b>",

            styles["Heading2"]

        )

    )

    doc.build(elements)

    return response

def bulk_delete_orders(request):

    if request.method == "POST":

        selected_orders = request.POST.getlist(
            "selected_orders"
        )

        deleted_count = Order.objects.filter(
            id__in=selected_orders
        ).count()

        Order.objects.filter(
            id__in=selected_orders
        ).delete()

        messages.success(
            request,
            f"{deleted_count} order(s) deleted successfully."
        )

    return redirect("dashboard:order_list")

def order_delete(request, pk):

    order = get_object_or_404(Order, pk=pk)

    if request.method == "POST":

        order_number = order.order_number

        order.delete()

        messages.success(
            request,
            f'Order "{order_number}" deleted successfully.'
        )

    return redirect("dashboard:order_list")

def product_list(request):

    search = request.GET.get("search", "")

    products = (
        Product.objects
        .select_related("category")
        .order_by("-created_at")
    )

    if search:
        products = products.filter(
            Q(name__icontains=search) |
            Q(category__name__icontains=search)
        )

    paginator = Paginator(products, 10)

    page = request.GET.get("page")

    products = paginator.get_page(page)

    context = {
        "products": products,
        "search": search,
    }

    return render(
        request,
        "dashboard/products/product_list.html",
        context,
    )

def product_create(request):

    if request.method == "POST":

        form = ProductForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            product = form.save(commit=False)

            # Automatically calculate discount percentage
            if (
                product.offer_price
                and product.offer_price < product.selling_price
            ):
                product.discount_percentage = int(
                    (
                        (product.selling_price - product.offer_price)
                        / product.selling_price
                    ) * 100
                ) 
            else:
                product.discount_percentage = 0

            product.save()

            gallery_images = request.FILES.getlist("gallery_images")

            for image in gallery_images:

                ProductImage.objects.create(
                    product=product,
                    image=image
                )

            return redirect("dashboard:product_list")

    else:

        form = ProductForm()

    return render(
        request,
        "dashboard/products/product_create.html",
        {
            "form": form
        }
    )

def product_update(request, pk):

    product = get_object_or_404(Product, pk=pk)

    gallery_images = ProductImage.objects.filter(product=product)


    if request.method == "POST":

        form = ProductForm(
            request.POST,
            request.FILES,
            instance=product
        )

        if form.is_valid():

            product = form.save()

            uploaded_images = request.FILES.getlist("gallery_images")

            for image in uploaded_images:

                ProductImage.objects.create(
                    product=product,
                    image=image
                )

            messages.success(
                request,
                "Product updated successfully."
            )

            return redirect("dashboard:product_list")

    else:

        form = ProductForm(instance=product)

    return render(
        request,
        "dashboard/products/product_update.html",
        {
            "form": form,
            "product": product,
            "gallery_images": gallery_images,
        }
    )

def delete_gallery_image(request, pk):

    image = get_object_or_404(ProductImage, pk=pk)

    product_id = image.product.id

    image.delete()

    messages.success(
        request,
        "Gallery image deleted successfully."
    )

    return redirect(
        "dashboard:product_update",
        pk=product_id
    )

def product_delete(request, pk):

    product = get_object_or_404(Product, pk=pk)

    if request.method == "POST":

        product_name = product.name

        product.delete()

        messages.success(
            request,
            f'"{product_name}" has been deleted successfully.'
        )

    return redirect("dashboard:product_list")

def category_list(request):

    search = request.GET.get("search", "")

    categories = get_category_tree()

    if search:

        categories = [
            category
            for category in categories
            if search.lower() in category.name.lower()
        ]

    paginator = Paginator(categories, 10)

    page_number = request.GET.get("page")

    page_obj = paginator.get_page(page_number)

    return render(
        request,
        "dashboard/category/category_list.html",
        {
            "page_obj": page_obj,
            "search": search,
        }
    )

def category_create(request):

    if request.method == "POST":

        form = CategoryForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Category added successfully."
            )

            return redirect("dashboard:category_list")

    else:

        form = CategoryForm()

    return render(
        request,
        "dashboard/category/category_create.html",
        {
            "form": form,
        }
    )

def category_update(request, pk):

    category = get_object_or_404(Category, pk=pk)

    if request.method == "POST":

        form = CategoryForm(
            request.POST,
            request.FILES,
            instance=category
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Category updated successfully."
            )

            return redirect("dashboard:category_list")

    else:

        form = CategoryForm(instance=category)

    return render(
        request,
        "dashboard/category/category_update.html",
        {
            "form": form,
            "category": category,
        }
    )

def category_delete(request, pk):

    category = get_object_or_404(Category, pk=pk)

    if request.method == "POST":

        # Check for child categories
        if category.subcategories.exists():

            messages.error(
                request,
                "Cannot delete this category because it has subcategories."
            )

            return redirect("dashboard:category_list")

        # Check for assigned products
        if category.products.exists():

            messages.error(
                request,
                "Cannot delete this category because products are assigned to it."
            )

            return redirect("dashboard:category_list")

        category_name = category.name

        category.delete()

        messages.success(
            request,
            f'"{category_name}" deleted successfully.'
        )

    return redirect("dashboard:category_list")

def customer_list(request):

    search = request.GET.get("search", "").strip()

    status = request.GET.get("status", "")

    date_filter = request.GET.get("date", "")


    # ==========================================
    # CUSTOMER QUERY
    # ==========================================

    customers = (
        User.objects
        .filter(is_staff=False)
        .annotate(
            total_orders=Count("order"),
            total_spent=Sum("order__total")
        )
        .order_by("-date_joined")
    )


    # ==========================================
    # SEARCH
    # ==========================================

    if search:

        customers = customers.filter(

            Q(username__icontains=search) |

            Q(email__icontains=search) |

            Q(phone_number__icontains=search)

        )


    # ==========================================
    # STATUS FILTER
    # ==========================================

    if status == "active":

        customers = customers.filter(
            is_active=True
        )

    elif status == "blocked":

        customers = customers.filter(
            is_active=False
        )


    # ==========================================
    # DATE FILTER
    # ==========================================

    today = timezone.now().date()


    if date_filter == "today":

        customers = customers.filter(
            date_joined__date=today
        )


    elif date_filter == "week":

        customers = customers.filter(
            date_joined__date__gte=
            today - timedelta(days=7)
        )


    elif date_filter == "month":

        customers = customers.filter(
            date_joined__month=today.month,
            date_joined__year=today.year
        )


    elif date_filter == "year":

        customers = customers.filter(
            date_joined__year=today.year
        )


    # ==========================================
    # PAGINATION
    # ==========================================

    paginator = Paginator(
        customers,
        10
    )

    page = request.GET.get("page")

    customers = paginator.get_page(page)


    # ==========================================
    # STATISTICS
    # ==========================================

    total_customers = User.objects.filter(
        is_staff=False
    ).count()


    active_customers = User.objects.filter(
        is_staff=False,
        is_active=True
    ).count()


    blocked_customers = User.objects.filter(
        is_staff=False,
        is_active=False
    ).count()


    new_customers = User.objects.filter(
        is_staff=False,
        date_joined__month=timezone.now().month,
        date_joined__year=timezone.now().year,
    ).count()


    # ==========================================
    # CONTEXT
    # ==========================================

    return render(
        request,
        "dashboard/customers/customer_list.html",
        {
            "customers": customers,

            "search": search,

            "status": status,

            "date_filter": date_filter,

            "total_customers": total_customers,

            "active_customers": active_customers,

            "blocked_customers": blocked_customers,

            "new_customers": new_customers,
        },
    )

def customer_detail(request, pk):

    customer = get_object_or_404(
        User,
        pk=pk,
        is_staff=False
    )

    orders = (
        Order.objects
        .filter(user=customer)
        .order_by("-created_at")
    )

    total_orders = orders.count()

    total_spent = (
        orders
        .filter(payment_status=True)
        .aggregate(
            total=Sum("total")
        )["total"] or 0
    )

    average_order = (
        orders
        .aggregate(
            avg=Avg("total")
        )["avg"] or 0
    )

    last_order = orders.first()

    context = {
        "customer": customer,
        "orders": orders,
        "total_orders": total_orders,
        "total_spent": total_spent,
        "average_order": average_order,
        "last_order": last_order,
    }

    return render(
        request,
        "dashboard/customers/customer_detail.html",
        context
    )
def customer_toggle_status(request, pk):

    customer = get_object_or_404(User, pk=pk)

    if request.method == "POST":

        customer.is_active = not customer.is_active
        customer.save()

        if customer.is_active:
            messages.success(
                request,
                f"{customer.username} has been unblocked successfully."
            )
        else:
            messages.success(
                request,
                f"{customer.username} has been blocked successfully."
            )

    return redirect(
        "dashboard:customer_detail",
        pk=pk
    )

def coupon_list(request):

    search = request.GET.get("search", "").strip()

    status = request.GET.get("status", "")

    coupons = Coupon.objects.all().order_by("-created_at")


    # ==========================================
    # SEARCH
    # ==========================================

    if search:

        coupons = coupons.filter(
            Q(code__icontains=search)
        )


    # ==========================================
    # STATUS FILTER
    # ==========================================

    now = timezone.now()

    if status == "active":

        coupons = coupons.filter(
            is_active=True,
            valid_from__lte=now,
            valid_until__gte=now
        )

    elif status == "inactive":

        coupons = coupons.filter(
            is_active=False
        )

    elif status == "expired":

        coupons = coupons.filter(
            valid_until__lt=now
        )


    # ==========================================
    # PAGINATION
    # ==========================================

    paginator = Paginator(
        coupons,
        10
    )

    page = request.GET.get("page")

    coupons = paginator.get_page(page)


    # ==========================================
    # STATISTICS
    # ==========================================

    total_coupons = Coupon.objects.count()


    active_coupons = Coupon.objects.filter(
        is_active=True,
        valid_from__lte=now,
        valid_until__gte=now
    ).count()


    expired_coupons = Coupon.objects.filter(
        valid_until__lt=now
    ).count()


    total_uses = (
        Coupon.objects.aggregate(
            total=Sum("used_count")
        )["total"] or 0
    )


    return render(
        request,
        "dashboard/coupons/coupon_list.html",
        {
            "coupons": coupons,

            "search": search,

            "status": status,

            "total_coupons": total_coupons,

            "active_coupons": active_coupons,

            "expired_coupons": expired_coupons,

            "total_uses": total_uses,
        }
    )

def create_coupon(request):

    if request.method == "POST":

        code = request.POST.get("code", "").strip().upper()

        discount_type = request.POST.get(
            "discount_type",
            "percentage"
        )

        discount_value = request.POST.get(
            "discount_value",
            ""
        ).strip()

        minimum_order_amount = request.POST.get(
            "minimum_order_amount",
            "0"
        ).strip()

        maximum_discount_amount = request.POST.get(
            "maximum_discount_amount",
            ""
        ).strip()

        usage_limit = request.POST.get(
            "usage_limit",
            ""
        ).strip()

        valid_from = request.POST.get(
            "valid_from",
            ""
        ).strip()

        valid_until = request.POST.get(
            "valid_until",
            ""
        ).strip()

        is_active = request.POST.get(
            "is_active"
        ) == "on"


        # ==========================================
        # COUPON CODE
        # ==========================================

        if not code:

            messages.error(
                request,
                "Coupon code is required."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        if Coupon.objects.filter(
            code__iexact=code
        ).exists():

            messages.error(
                request,
                "A coupon with this code already exists."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        # ==========================================
        # DISCOUNT TYPE
        # ==========================================

        if discount_type not in [
            "percentage",
            "fixed"
        ]:

            messages.error(
                request,
                "Invalid discount type."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        # ==========================================
        # DECIMAL VALUES
        # ==========================================

        try:

            discount_value = Decimal(
                discount_value
            )

            minimum_order_amount = Decimal(
                minimum_order_amount or "0"
            )

            if maximum_discount_amount:

                maximum_discount_amount = Decimal(
                    maximum_discount_amount
                )

            else:

                maximum_discount_amount = None

        except InvalidOperation:

            messages.error(
                request,
                "Please enter valid discount amounts."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        # ==========================================
        # DISCOUNT VALIDATION
        # ==========================================

        if discount_value <= 0:

            messages.error(
                request,
                "Discount value must be greater than zero."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        if discount_type == "percentage":

            if discount_value > 100:

                messages.error(
                    request,
                    "Percentage discount cannot exceed 100%."
                )

                return redirect(
                    "dashboard:create_coupon"
                )

        else:

            # Fixed discount cannot be greater than
            # the minimum order amount when one exists.

            if (
                minimum_order_amount > 0
                and discount_value > minimum_order_amount
            ):

                messages.error(
                    request,
                    "Fixed discount cannot be greater than the minimum order amount."
                )

                return redirect(
                    "dashboard:create_coupon"
                )


        # ==========================================
        # MINIMUM ORDER
        # ==========================================

        if minimum_order_amount < 0:

            messages.error(
                request,
                "Minimum order amount cannot be negative."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        # ==========================================
        # MAXIMUM DISCOUNT
        # ==========================================

        if maximum_discount_amount is not None:

            if maximum_discount_amount <= 0:

                messages.error(
                    request,
                    "Maximum discount must be greater than zero."
                )

                return redirect(
                    "dashboard:create_coupon"
                )


            if discount_type == "fixed":

                messages.error(
                    request,
                    "Maximum discount is only applicable to percentage coupons."
                )

                return redirect(
                    "dashboard:create_coupon"
                )


        # ==========================================
        # USAGE LIMIT
        # ==========================================

        if usage_limit:

            try:

                usage_limit = int(
                    usage_limit
                )

            except ValueError:

                messages.error(
                    request,
                    "Usage limit must be a valid number."
                )

                return redirect(
                    "dashboard:create_coupon"
                )


            if usage_limit <= 0:

                messages.error(
                    request,
                    "Usage limit must be greater than zero."
                )

                return redirect(
                    "dashboard:create_coupon"
                )

        else:

            usage_limit = None


        # ==========================================
        # DATE PARSING
        # ==========================================

        if not valid_from or not valid_until:

            messages.error(
                request,
                "Please provide both start and expiry dates."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        try:

            valid_from = datetime.fromisoformat(
                valid_from
            )

            valid_until = datetime.fromisoformat(
                valid_until
            )

        except ValueError:

            messages.error(
                request,
                "Invalid date format."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        # ==========================================
        # TIMEZONE
        # ==========================================

        if timezone.is_naive(valid_from):

            valid_from = timezone.make_aware(
                valid_from
            )


        if timezone.is_naive(valid_until):

            valid_until = timezone.make_aware(
                valid_until
            )


        # ==========================================
        # DATE ORDER
        # ==========================================

        if valid_until <= valid_from:

            messages.error(
                request,
                "Expiry date must be later than the start date."
            )

            return redirect(
                "dashboard:create_coupon"
            )


        # ==========================================
        # CREATE COUPON
        # ==========================================

        Coupon.objects.create(

            code=code,

            discount_type=discount_type,

            discount_value=discount_value,

            minimum_order_amount=
                minimum_order_amount,

            maximum_discount_amount=
                maximum_discount_amount,

            usage_limit=usage_limit,

            valid_from=valid_from,

            valid_until=valid_until,

            is_active=is_active,

        )


        messages.success(
            request,
            f"Coupon {code} created successfully."
        )


        return redirect(
            "dashboard:coupon_list"
        )


    return render(
        request,
        "dashboard/coupons/coupon_create.html"
    )

def edit_coupon(request, coupon_id):

    coupon = get_object_or_404(
        Coupon,
        id=coupon_id
    )

    if request.method == "POST":

        code = request.POST.get(
            "code",
            ""
        ).strip().upper()

        discount_type = request.POST.get(
            "discount_type",
            "percentage"
        )

        discount_value = request.POST.get(
            "discount_value",
            ""
        ).strip()

        minimum_order_amount = request.POST.get(
            "minimum_order_amount",
            "0"
        ).strip()

        maximum_discount_amount = request.POST.get(
            "maximum_discount_amount",
            ""
        ).strip()

        usage_limit = request.POST.get(
            "usage_limit",
            ""
        ).strip()

        valid_from = request.POST.get(
            "valid_from",
            ""
        ).strip()

        valid_until = request.POST.get(
            "valid_until",
            ""
        ).strip()

        is_active = request.POST.get(
            "is_active"
        ) == "on"


        # ------------------------------
        # CODE
        # ------------------------------

        if not code:

            messages.error(
                request,
                "Coupon code is required."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        if Coupon.objects.filter(
            code__iexact=code
        ).exclude(
            id=coupon.id
        ).exists():

            messages.error(
                request,
                "Another coupon already uses this code."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        # ------------------------------
        # DISCOUNT TYPE
        # ------------------------------

        if discount_type not in [
            "percentage",
            "fixed"
        ]:

            messages.error(
                request,
                "Invalid discount type."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        # ------------------------------
        # DECIMAL VALUES
        # ------------------------------

        try:

            discount_value = Decimal(
                discount_value
            )

            minimum_order_amount = Decimal(
                minimum_order_amount or "0"
            )

            if maximum_discount_amount:

                maximum_discount_amount = Decimal(
                    maximum_discount_amount
                )

            else:

                maximum_discount_amount = None

        except InvalidOperation:

            messages.error(
                request,
                "Please enter valid discount amounts."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        # ------------------------------
        # DISCOUNT VALIDATION
        # ------------------------------

        if discount_value <= 0:

            messages.error(
                request,
                "Discount value must be greater than zero."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        if discount_type == "percentage":

            if discount_value > 100:

                messages.error(
                    request,
                    "Percentage discount cannot exceed 100%."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )

        else:

            if (
                minimum_order_amount > 0
                and discount_value > minimum_order_amount
            ):

                messages.error(
                    request,
                    "Fixed discount cannot be greater than the minimum order amount."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )


        # ------------------------------
        # MINIMUM ORDER
        # ------------------------------

        if minimum_order_amount < 0:

            messages.error(
                request,
                "Minimum order amount cannot be negative."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        # ------------------------------
        # MAXIMUM DISCOUNT
        # ------------------------------

        if maximum_discount_amount is not None:

            if maximum_discount_amount <= 0:

                messages.error(
                    request,
                    "Maximum discount must be greater than zero."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )


            if discount_type == "fixed":

                messages.error(
                    request,
                    "Maximum discount is only applicable to percentage coupons."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )


        # ------------------------------
        # USAGE LIMIT
        # ------------------------------

        if usage_limit:

            try:

                usage_limit = int(
                    usage_limit
                )

            except ValueError:

                messages.error(
                    request,
                    "Usage limit must be a valid number."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )


            if usage_limit <= 0:

                messages.error(
                    request,
                    "Usage limit must be greater than zero."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )

            if usage_limit < coupon.used_count:

                messages.error(
                    request,
                    f"Usage limit cannot be less than the current usage ({coupon.used_count})."
                )

                return redirect(
                    "dashboard:edit_coupon",
                    coupon_id=coupon.id
                )

        else:

            usage_limit = None


        # ------------------------------
        # DATES
        # ------------------------------

        if not valid_from or not valid_until:

            messages.error(
                request,
                "Please provide both start and expiry dates."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        try:

            valid_from = datetime.fromisoformat(
                valid_from
            )

            valid_until = datetime.fromisoformat(
                valid_until
            )

        except ValueError:

            messages.error(
                request,
                "Invalid date format."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        if timezone.is_naive(valid_from):

            valid_from = timezone.make_aware(
                valid_from
            )


        if timezone.is_naive(valid_until):

            valid_until = timezone.make_aware(
                valid_until
            )


        if valid_until <= valid_from:

            messages.error(
                request,
                "Expiry date must be later than the start date."
            )

            return redirect(
                "dashboard:edit_coupon",
                coupon_id=coupon.id
            )


        # ------------------------------
        # UPDATE
        # ------------------------------

        coupon.code = code

        coupon.discount_type = discount_type

        coupon.discount_value = discount_value

        coupon.minimum_order_amount = (
            minimum_order_amount
        )

        coupon.maximum_discount_amount = (
            maximum_discount_amount
        )

        coupon.usage_limit = usage_limit

        coupon.valid_from = valid_from

        coupon.valid_until = valid_until

        coupon.is_active = is_active

        coupon.save()


        messages.success(
            request,
            f"Coupon {code} updated successfully."
        )


        return redirect(
            "dashboard:coupon_list"
        )


    return render(
        request,
        "dashboard/coupons/coupon_create.html",
        {
            "coupon": coupon,
            "edit_mode": True,
        }
    )

def coupon_detail(request, coupon_id):

    coupon = get_object_or_404(
        Coupon,
        id=coupon_id
    )

    now = timezone.now()

    if not coupon.is_active:

        coupon_status = "inactive"
        coupon_status_display = "Inactive"

    elif coupon.valid_until < now:

        coupon_status = "expired"
        coupon_status_display = "Expired"

    elif coupon.valid_from > now:

        coupon_status = "upcoming"
        coupon_status_display = "Upcoming"

    else:

        coupon_status = "active"
        coupon_status_display = "Active"

    return render(
        request,
        "dashboard/coupons/coupon_detail.html",
        {
            "coupon": coupon,
            "coupon_status": coupon_status,
            "coupon_status_display": coupon_status_display,
        }
    )

def delete_coupon(request, coupon_id):

    if request.method == "POST":

        coupon = get_object_or_404(
            Coupon,
            id=coupon_id
        )

        coupon_code = coupon.code

        coupon.delete()

        messages.success(
            request,
            f"Coupon '{coupon_code}' deleted permanently."
        )

    return redirect("dashboard:coupon_list")  

def report(request):

    date_filter = request.GET.get("date", "all")
    today = timezone.now().date()

    orders = Order.objects.all()

    # -------------------------
    # DATE FILTER
    # -------------------------

    if date_filter == "today":
        orders = orders.filter(
            created_at__date=today
        )

    elif date_filter == "week":
        orders = orders.filter(
            created_at__date__gte=today - timedelta(days=7)
        )

    elif date_filter == "month":
        orders = orders.filter(
            created_at__month=today.month,
            created_at__year=today.year
        )

    elif date_filter == "year":
        orders = orders.filter(
            created_at__year=today.year
        )

    # -------------------------
    # BASIC REPORT DATA
    # -------------------------

    total_orders = orders.count()

    paid_delivered_orders = orders.filter(
        payment_status=True,
        status="delivered"
    )

    total_revenue = (
        paid_delivered_orders.aggregate(
            total=Sum("total")
        )["total"] or 0
    )

    average_order_value = (
        paid_delivered_orders.aggregate(
            average=Avg("total")
        )["average"] or 0
    )

    pending_orders = orders.filter(
        status="pending"
    ).count()

    confirmed_orders = orders.filter(
        status="confirmed"
    ).count()

    shipped_orders = orders.filter(
        status="shipped"
    ).count()

    delivered_orders = orders.filter(
        status="delivered"
    ).count()

    cancelled_orders = orders.filter(
        status="cancelled"
    ).count()

    cod_orders = orders.filter(
        payment_method="cod"
    ).count()

    razorpay_orders = orders.filter(
        payment_method="razorpay"
    ).count()

    total_customers = User.objects.filter(
        is_staff=False
    ).count()

    total_products = Product.objects.count()

    # -------------------------
    # SALES CHART DATA
    # -------------------------

    chart_orders = paid_delivered_orders

    if date_filter == "year":

        sales_data = (
            chart_orders
            .annotate(period=TruncMonth("created_at"))
            .values("period")
            .annotate(
                revenue=Sum("total"),
                order_count=Count("id")
            )
            .order_by("period")
        )

        chart_labels = [
            item["period"].strftime("%b %Y")
            for item in sales_data
        ]

    else:

        sales_data = (
            chart_orders
            .annotate(period=TruncDate("created_at"))
            .values("period")
            .annotate(
                revenue=Sum("total"),
                order_count=Count("id")
            )
            .order_by("period")
        )

        chart_labels = [
            item["period"].strftime("%d %b")
            for item in sales_data
        ]

    chart_revenue = [
        float(item["revenue"])
        for item in sales_data
    ]

    chart_orders_count = [
        item["order_count"]
        for item in sales_data
    ]

    # -------------------------
    # BEST-SELLING PRODUCTS
    # -------------------------

    best_selling_products = (
        OrderItem.objects
        .filter(
            order__in=paid_delivered_orders
        )
        .values(
            "product__name"
        )
        .annotate(
            units_sold=Sum("quantity")
        )
        .order_by("-units_sold")[:10]
    )

    # -------------------------
    # CONTEXT
    # -------------------------

    context = {
        "date_filter": date_filter,

        "total_orders": total_orders,
        "total_revenue": total_revenue,
        "average_order_value": average_order_value,

        "pending_orders": pending_orders,
        "confirmed_orders": confirmed_orders,
        "shipped_orders": shipped_orders,
        "delivered_orders": delivered_orders,
        "cancelled_orders": cancelled_orders,

        "cod_orders": cod_orders,
        "razorpay_orders": razorpay_orders,

        "total_customers": total_customers,
        "total_products": total_products,

        # Chart data
        "chart_labels": json.dumps(chart_labels),
        "chart_revenue": json.dumps(chart_revenue),
        "chart_orders": json.dumps(chart_orders_count),

        "best_selling_products": best_selling_products,
    }

    return render(
        request,
        "dashboard/reports/report.html",
        context
    )

def download_sales_report_pdf(request):

    date_filter = request.GET.get("date", "all")
    today = timezone.now().date()

    orders = Order.objects.all()

    # -------------------------
    # DATE FILTER
    # -------------------------

    if date_filter == "today":
        orders = orders.filter(
            created_at__date=today
        )

    elif date_filter == "week":
        orders = orders.filter(
            created_at__date__gte=today - timedelta(days=7)
        )

    elif date_filter == "month":
        orders = orders.filter(
            created_at__month=today.month,
            created_at__year=today.year
        )

    elif date_filter == "year":
        orders = orders.filter(
            created_at__year=today.year
        )

    # -------------------------
    # REPORT DATA
    # -------------------------

    paid_delivered_orders = orders.filter(
        payment_status=True,
        status="delivered"
    )

    total_orders = orders.count()

    total_revenue = (
        paid_delivered_orders.aggregate(
            total=Sum("total")
        )["total"] or 0
    )

    average_order_value = (
        paid_delivered_orders.aggregate(
            average=Avg("total")
        )["average"] or 0
    )

    pending_orders = orders.filter(
        status="pending"
    ).count()

    confirmed_orders = orders.filter(
        status="confirmed"
    ).count()

    shipped_orders = orders.filter(
        status="shipped"
    ).count()

    delivered_orders = orders.filter(
        status="delivered"
    ).count()

    cancelled_orders = orders.filter(
        status="cancelled"
    ).count()

    cod_orders = orders.filter(
        payment_method="cod"
    ).count()

    razorpay_orders = orders.filter(
        payment_method="razorpay"
    ).count()

    best_selling_products = (
        OrderItem.objects
        .filter(order__in=paid_delivered_orders)
        .values("product__name")
        .annotate(
            units_sold=Sum("quantity")
        )
        .order_by("-units_sold")[:10]
    )

    # -------------------------
    # PERIOD NAME
    # -------------------------

    period_names = {
        "all": "All Time",
        "today": "Today",
        "week": "Last 7 Days",
        "month": "This Month",
        "year": "This Year",
    }

    period_name = period_names.get(
        date_filter,
        "All Time"
    )

    # -------------------------
    # PDF RESPONSE
    # -------------------------

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="sales_report.pdf"'
    )

    doc = SimpleDocTemplate(
        response,
        rightMargin=35,
        leftMargin=35,
        topMargin=35,
        bottomMargin=35,
    )

    styles = getSampleStyleSheet()

    elements = []

    # -------------------------
    # TITLE
    # -------------------------

    elements.append(
        Paragraph(
            "Sales Report",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            f"Report Period: {period_name}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 20)
    )

    # -------------------------
    # SUMMARY
    # -------------------------

    summary_data = [
        ["Metric", "Value"],
        ["Total Orders", str(total_orders)],
        ["Total Revenue", f"₹{total_revenue}"],
        [
            "Average Order Value",
            f"₹{average_order_value:.2f}"
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[250, 220]
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.orange
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "BACKGROUND",
                (0, 1),
                (-1, -1),
                colors.whitesmoke
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    elements.append(summary_table)

    elements.append(
        Spacer(1, 20)
    )

    # -------------------------
    # ORDER STATUS
    # -------------------------

    elements.append(
        Paragraph(
            "Order Status",
            styles["Heading2"]
        )
    )

    status_data = [
        ["Status", "Orders"],
        ["Pending", str(pending_orders)],
        ["Confirmed", str(confirmed_orders)],
        ["Shipped", str(shipped_orders)],
        ["Delivered", str(delivered_orders)],
        ["Cancelled", str(cancelled_orders)],
    ]

    status_table = Table(
        status_data,
        colWidths=[250, 220]
    )

    status_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.orange
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    elements.append(status_table)

    elements.append(
        Spacer(1, 20)
    )

    # -------------------------
    # PAYMENT METHODS
    # -------------------------

    elements.append(
        Paragraph(
            "Payment Methods",
            styles["Heading2"]
        )
    )

    payment_data = [
        ["Payment Method", "Orders"],
        ["Cash on Delivery", str(cod_orders)],
        ["Razorpay", str(razorpay_orders)],
    ]

    payment_table = Table(
        payment_data,
        colWidths=[250, 220]
    )

    payment_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.orange
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    elements.append(payment_table)

    elements.append(
        Spacer(1, 20)
    )

    # -------------------------
    # BEST SELLING PRODUCTS
    # -------------------------

    elements.append(
        Paragraph(
            "Best-Selling Products",
            styles["Heading2"]
        )
    )

    product_data = [
        ["#", "Product", "Units Sold"]
    ]

    for index, product in enumerate(
        best_selling_products,
        start=1
    ):
        product_data.append([
            str(index),
            product["product__name"],
            str(product["units_sold"]),
        ])

    if len(product_data) == 1:
        product_data.append([
            "-",
            "No sales data available",
            "-"
        ])

    product_table = Table(
        product_data,
        colWidths=[40, 300, 130]
    )

    product_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.orange
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )

    elements.append(product_table)

    elements.append(
        Spacer(1, 25)
    )

    elements.append(
        Paragraph(
            "Generated from Casaio Admin Dashboard",
            styles["Normal"]
        )
    )

    doc.build(elements)

    return response

def download_sales_report_excel(request):

    date_filter = request.GET.get("date", "all")
    today = timezone.now().date()

    orders = Order.objects.all()

    # -------------------------
    # DATE FILTER
    # -------------------------

    if date_filter == "today":
        orders = orders.filter(
            created_at__date=today
        )

    elif date_filter == "week":
        orders = orders.filter(
            created_at__date__gte=today - timedelta(days=7)
        )

    elif date_filter == "month":
        orders = orders.filter(
            created_at__month=today.month,
            created_at__year=today.year
        )

    elif date_filter == "year":
        orders = orders.filter(
            created_at__year=today.year
        )

    # -------------------------
    # REPORT DATA
    # -------------------------

    paid_delivered_orders = orders.filter(
        payment_status=True,
        status="delivered"
    )

    total_orders = orders.count()

    total_revenue = (
        paid_delivered_orders.aggregate(
            total=Sum("total")
        )["total"] or 0
    )

    average_order_value = (
        paid_delivered_orders.aggregate(
            average=Avg("total")
        )["average"] or 0
    )

    pending_orders = orders.filter(
        status="pending"
    ).count()

    confirmed_orders = orders.filter(
        status="confirmed"
    ).count()

    shipped_orders = orders.filter(
        status="shipped"
    ).count()

    delivered_orders = orders.filter(
        status="delivered"
    ).count()

    cancelled_orders = orders.filter(
        status="cancelled"
    ).count()

    cod_orders = orders.filter(
        payment_method="cod"
    ).count()

    razorpay_orders = orders.filter(
        payment_method="razorpay"
    ).count()

    best_selling_products = (
        OrderItem.objects
        .filter(order__in=paid_delivered_orders)
        .values("product__name")
        .annotate(
            units_sold=Sum("quantity")
        )
        .order_by("-units_sold")[:10]
    )

    # -------------------------
    # EXCEL WORKBOOK
    # -------------------------

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Sales Report"

    # -------------------------
    # TITLE
    # -------------------------

    worksheet["A1"] = "Casaio Sales Report"
    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = (
        f"Report Period: {date_filter.replace('_', ' ').title()}"
    )

    # -------------------------
    # SUMMARY
    # -------------------------

    worksheet["A4"] = "Summary"
    worksheet["A4"].font = Font(
        bold=True,
        size=13
    )

    summary_headers = [
        "Metric",
        "Value"
    ]

    for column, value in enumerate(
        summary_headers,
        start=1
    ):
        cell = worksheet.cell(
            row=5,
            column=column,
            value=value
        )
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="FF6B00"
        )

    summary_rows = [
        ["Total Orders", total_orders],
        ["Total Revenue", float(total_revenue)],
        [
            "Average Order Value",
            float(average_order_value)
        ],
    ]

    for row_index, row in enumerate(
        summary_rows,
        start=6
    ):
        for column_index, value in enumerate(
            row,
            start=1
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

    # -------------------------
    # ORDER STATUS
    # -------------------------

    status_start = 11

    worksheet.cell(
        row=status_start,
        column=1,
        value="Order Status"
    ).font = Font(
        bold=True,
        size=13
    )

    status_headers = [
        "Status",
        "Orders"
    ]

    for column, value in enumerate(
        status_headers,
        start=1
    ):
        cell = worksheet.cell(
            row=status_start + 1,
            column=column,
            value=value
        )
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="FF6B00"
        )

    status_rows = [
        ["Pending", pending_orders],
        ["Confirmed", confirmed_orders],
        ["Shipped", shipped_orders],
        ["Delivered", delivered_orders],
        ["Cancelled", cancelled_orders],
    ]

    for row_index, row in enumerate(
        status_rows,
        start=status_start + 2
    ):
        for column_index, value in enumerate(
            row,
            start=1
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

    # -------------------------
    # PAYMENT METHODS
    # -------------------------

    payment_start = 20

    worksheet.cell(
        row=payment_start,
        column=1,
        value="Payment Methods"
    ).font = Font(
        bold=True,
        size=13
    )

    payment_headers = [
        "Payment Method",
        "Orders"
    ]

    for column, value in enumerate(
        payment_headers,
        start=1
    ):
        cell = worksheet.cell(
            row=payment_start + 1,
            column=column,
            value=value
        )
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="FF6B00"
        )

    payment_rows = [
        ["Cash on Delivery", cod_orders],
        ["Razorpay", razorpay_orders],
    ]

    for row_index, row in enumerate(
        payment_rows,
        start=payment_start + 2
    ):
        for column_index, value in enumerate(
            row,
            start=1
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

    # -------------------------
    # BEST SELLING PRODUCTS
    # -------------------------

    product_start = 26

    worksheet.cell(
        row=product_start,
        column=1,
        value="Best-Selling Products"
    ).font = Font(
        bold=True,
        size=13
    )

    product_headers = [
        "Rank",
        "Product",
        "Units Sold"
    ]

    for column, value in enumerate(
        product_headers,
        start=1
    ):
        cell = worksheet.cell(
            row=product_start + 1,
            column=column,
            value=value
        )
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="FF6B00"
        )

    for index, product in enumerate(
        best_selling_products,
        start=1
    ):
        worksheet.cell(
            row=product_start + 1 + index,
            column=1,
            value=index
        )

        worksheet.cell(
            row=product_start + 1 + index,
            column=2,
            value=product["product__name"]
        )

        worksheet.cell(
            row=product_start + 1 + index,
            column=3,
            value=product["units_sold"]
        )

    # -------------------------
    # COLUMN WIDTHS
    # -------------------------

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 18

    # -------------------------
    # RESPONSE
    # -------------------------

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="sales_report.xlsx"'
    )

    workbook.save(response)

    return response